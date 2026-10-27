"""Build a cleaned Opening Stock proposal + report from the template's Opening Stock tab.
Usage: python3 scripts/clean_opening_stock.py <template.xlsx> <outdir>
Rule set (for approval):
  A. Group rows by (Warehouse, Item Code, Batch Number).
  B. Net quantity and value inside each group (a -2 row usually cancels a +2 row that was keyed with a different expiry).
  C. Net quantity <= 0 -> excluded. Net > 0 -> one merged row; expiry = earliest expiry among the positive rows.
  D. Blank Stock Value is NOT invented: stays blank (valuation treats blank as 0) and is listed in the report.
Also reports what the literal rule 'drop negative rows, then merge' would have produced, for comparison.
"""
import sys, collections, os, openpyxl
from openpyxl.styles import Font, PatternFill
src, out = sys.argv[1], sys.argv[2]
os.makedirs(out, exist_ok=True)
ws = openpyxl.load_workbook(src, data_only=True)["Opening Stock"]
H = [c.value for c in ws[1]]
rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
g = collections.OrderedDict()
for r in rows: g.setdefault((r[0], r[1], r[4]), []).append(r)
cleaned, removed_neg, merged, blank_val = [], [], [], []
lit = collections.defaultdict(float)
for k, rs in g.items():
    q = sum(x[5] or 0 for x in rs)
    vals = [x[8] for x in rs if x[8] is not None]
    v = round(sum(vals), 2) if vals else None
    pos = [x for x in rs if (x[5] or 0) > 0]
    for x in rs:
        if (x[5] or 0) > 0: lit[k[0]] += x[5]
    if len(rs) > 1: merged.append((k, rs, q, v))
    if q <= 0:
        removed_neg.append((k, rs, q)); continue
    base = pos[0]
    exps = sorted(x[6] for x in pos if x[6])
    row = [base[0], base[1], base[2], base[3] or next((x[3] for x in rs if x[3]), None), base[4], q, exps[0] if exps else None, base[7], v]
    cleaned.append(row)
    if v is None: blank_val.append(row)
hdr_fill = PatternFill("solid", fgColor="1F4E79")
def sheet(wb, name, head, data):
    s = wb.create_sheet(name); s.append(head)
    for c in s[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr_fill
    for d in data: s.append(d)
    s.freeze_panes = "A2"
    for i, h in enumerate(head, 1): s.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(14, min(40, len(str(h)) + 4))
wb = openpyxl.Workbook(); wb.remove(wb.active)
sheet(wb, "Opening Stock", H, cleaned)
wb.save(f"{out}/Opening_Stock_CLEANED_for_approval.xlsx")
by = lambda rs, w, i: sum(x[i] or 0 for x in rs if x[0] == w)
sumr = []
for w in sorted({r[0] for r in rows}):
    cr = [r for r in cleaned if r[0] == w]
    sumr.append([w, len([r for r in rows if r[0] == w]), by(rows, w, 5), round(by(rows, w, 8), 2), len(cr), sum(r[5] for r in cr), round(sum(r[8] or 0 for r in cr), 2), lit[w]])
rep = openpyxl.Workbook(); rep.remove(rep.active)
sheet(rep, "Summary", ["Warehouse", "Rows before", "Qty before (net)", "Value before", "Rows after", "Qty after", "Value after", "Qty if negatives simply dropped (literal rule)"], sumr)
sheet(rep, "Excluded (net<=0)", ["Warehouse", "Item Code", "Batch", "Rows in group", "Net Qty", "Detail"],
      [[k[0], k[1], k[2], len(rs), q, "; ".join(f"{x[5]}@{x[6]}" for x in rs)] for k, rs, q in removed_neg])
sheet(rep, "Merged duplicates", ["Warehouse", "Item Code", "Batch", "Rows merged", "Net Qty", "Net Value", "Detail (qty@expiry)", "Expiry differs?"],
      [[k[0], k[1], k[2], len(rs), q, v, "; ".join(f"{x[5]}@{x[6]}" for x in rs), "YES" if len({x[6] for x in rs}) > 1 else ""] for k, rs, q, v in merged])
sheet(rep, "Blank Stock Value", H, blank_val)
rep.save(f"{out}/Opening_Stock_cleaning_report.xlsx")
print("before", len(rows), "after", len(cleaned), "excluded groups", len(removed_neg), "merged groups", len(merged), "blank value", len(blank_val))
for r in sumr: print(r)
